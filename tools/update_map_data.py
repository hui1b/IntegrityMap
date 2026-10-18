#!/usr/bin/env python3
"""Refresh the map dataset from the project's source Excel file.

This script reads the new sheet, normalizes the rows, and writes a fresh
window.MAP_DATA payload to data.js plus a JSON copy to data.json.
"""

from __future__ import annotations

import json
import os
import re
import time
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import quote
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA_JS = ROOT / "data.js"
DATA_JSON = ROOT / "data.json"

CATEGORY_MAP = {
    "历史遗迹": "110865",
    "纪念场馆": "110866",
    "教育基地": "110867",
    "艺术作品": "110868",
    "文化品牌": "110869",
    "学科课程": "110870",
}


def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def find_source_file() -> Path:
    names = [
        p for p in sorted(os.listdir(ROOT))
        if p.lower().endswith((".xlsx", ".xls", ".csv"))
    ]
    if not names:
        raise FileNotFoundError("No spreadsheet found in project root. Please add 廉洁文化地图素材.xlsx or a CSV file.")
    return ROOT / names[0]


def read_rows(path: Path):
    if path.suffix.lower() == ".csv":
        import csv
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        return rows

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return list(ws.iter_rows(values_only=True))


def build_points(rows):
    points = []
    for row in rows[3:]:
        if not any(v is not None and str(v).strip() for v in row[:5]):
            continue
        idx, school, title, category, province = (row + (None,) * 5)[:5]
        school = normalize_text(school)
        title = normalize_text(title)
        category = normalize_text(category)
        province = normalize_text(province)
        if not school or not title:
            continue
        point = {
            "id": str(idx).strip() if idx is not None else str(len(points) + 1),
            "title": title,
            "schoolName": school,
            "category": category,
            "categoryId": CATEGORY_MAP.get(category, ""),
            "address": f"{school} {title} {province}".strip(),
            "summary": "",
            "link": "",
            "province": province,
            "lng": None,
            "lat": None,
        }
        points.append(point)
    return points


def geocode_address(address: str):
    text = normalize_text(address)
    if not text:
        return None

    query = quote(text)
    url = f"https://nominatim.openstreetmap.org/search?format=jsonv2&limit=1&polygon_geojson=0&addressdetails=1&q={query}"
    req = Request(
        url,
        headers={
            "User-Agent": "NationalIntegrityCultureMap/1.0 (+https://example.com)",
            "Accept-Language": "zh-CN,zh;q=0.9",
        },
    )
    try:
        with urlopen(req, timeout=20) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except Exception:
        return None

    if not payload:
        return None

    first = payload[0]
    try:
        lon = float(first["lon"])
        lat = float(first["lat"])
        return {"lng": lon, "lat": lat}
    except (KeyError, TypeError, ValueError):
        return None


def geocode_points(points):
    geocoded = 0
    for index, point in enumerate(points, start=1):
        if point.get("lng") is not None and point.get("lat") is not None:
            continue

        candidate = point.get("address") or f"{point.get('schoolName', '')} {point.get('title', '')} {point.get('province', '')}"
        coords = geocode_address(candidate)
        if coords:
            point["lng"] = coords["lng"]
            point["lat"] = coords["lat"]
            geocoded += 1
        time.sleep(0.35)

    return geocoded


def normalize_title(value: str) -> str:
    """Drop punctuation/whitespace so titles from different sources can be compared."""
    return re.sub(r"[\s《》（）()、·—\-–,，。.；;:：“”\"'‘’]", "", value or "")


def load_existing_points():
    """Read the details already present in data.json (if any)."""
    if not DATA_JSON.exists():
        return []
    try:
        obj = json.loads(DATA_JSON.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    return [p for p in obj.get("points", []) if isinstance(p, dict)]


def merge_details(points, existing):
    """Carry summary/link over from the previous dataset.

    Matching is school-scoped so a generic title such as "廉洁文化教育基地"
    cannot be pinned to another university's entry. Only empty fields are
    filled; anything already present in the new data wins.
    """
    by_school = defaultdict(list)
    for p in existing:
        if p.get("schoolName"):
            by_school[p["schoolName"]].append(p)

    filled = 0
    for point in points:
        if (point.get("summary") or "").strip() and (point.get("link") or "").strip():
            continue
        target_title = normalize_title(point.get("title", ""))
        candidates = by_school.get(point.get("schoolName", ""), [])

        match = next((c for c in candidates if c.get("title") == point.get("title")), None)
        if match is None:
            match = next((c for c in candidates if normalize_title(c.get("title", "")) == target_title), None)
        if match is None:
            contained = [
                (abs(len(normalize_title(c.get("title", ""))) - len(target_title)), c)
                for c in candidates
                if target_title
                and normalize_title(c.get("title", ""))
                and (
                    target_title in normalize_title(c.get("title", ""))
                    or normalize_title(c.get("title", "")) in target_title
                )
            ]
            if contained:
                contained.sort(key=lambda item: item[0])
                match = contained[0][1]
        if match is None:
            best, best_ratio = None, 0.0
            for c in candidates:
                ratio = SequenceMatcher(None, target_title, normalize_title(c.get("title", ""))).ratio()
                if ratio > best_ratio:
                    best_ratio, best = ratio, c
            if best is not None and best_ratio >= 0.80:
                match = best

        if match is None:
            continue

        if not (point.get("summary") or "").strip() and match.get("summary"):
            point["summary"] = match["summary"]
        if not (point.get("link") or "").strip() and match.get("link"):
            point["link"] = match["link"]
        if (point.get("summary") or "").strip() or (point.get("link") or "").strip():
            filled += 1

    return filled


def write_data(points):
    categories = [
        {"id": "", "name": "全部", "count": len(points)},
        {"id": "110865", "name": "历史遗迹", "count": 0},
        {"id": "110866", "name": "纪念场馆", "count": 0},
        {"id": "110867", "name": "教育基地", "count": 0},
        {"id": "110868", "name": "艺术作品", "count": 0},
        {"id": "110869", "name": "文化品牌", "count": 0},
        {"id": "110870", "name": "学科课程", "count": 0},
    ]
    counts = {c["id"]: c for c in categories if c["id"]}
    for p in points:
        cid = str(p.get("categoryId", "")).strip()
        if cid in counts:
            counts[cid]["count"] += 1

    meta = {
        "title": "廉洁文化地图",
        "subtitle": "全国高校廉洁文化点位",
        "source": "全国高校思想政治工作网",
        "totalPoints": len(points),
        "totalSchools": len({p.get("schoolName", "") for p in points if p.get("schoolName")}),
        "categories": categories,
    }
    obj = {"meta": meta, "points": points}
    DATA_JS.write_text("window.MAP_DATA = " + json.dumps(obj, ensure_ascii=False, indent=1) + ";\n", encoding="utf-8")
    DATA_JSON.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    src = find_source_file()
    rows = read_rows(src)
    if len(rows) < 4:
        raise SystemExit(f"No valid data rows found in {src.name}.")
    points = build_points(rows)
    geocoded_count = geocode_points(points)
    existing = load_existing_points()
    restored = merge_details(points, existing)
    write_data(points)
    print(
        f"Updated data from {src.name}: {len(points)} points; "
        f"geocoded {geocoded_count} points; restored details on {restored} points"
    )


if __name__ == "__main__":
    main()
